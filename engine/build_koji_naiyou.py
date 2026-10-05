# -*- coding: utf-8 -*-
"""（1-2）工事内容.xlsx を雛形から生成する。

使い方:
    python build_koji_naiyou.py 入力.json

入力JSONの仕様は ../scripts/README.md を参照。
雛形 templates/（1-2）工事内容_雛形.xlsx は「1明細＝10行」の均一ブロックが
並んだだけの様式で、罫線・結合・印刷設定はすべてブロック単位で同一。
本スクリプトは必要ブロック数だけ残して余りを削除し、各ブロックに値を書き込む。
"""
import json
import shutil
import sys
from pathlib import Path

import openpyxl

SHEET = "本工事費表"
FIRST_ROW = 4          # 先頭ブロックの基準行
BLOCK = 10             # 1ブロックの行数
LEVEL_COL = {"B": 2, "C": 3, "D": 4, "E": 5, "F": 6, "G": 7}
COL_SPEC = 7           # 規格（基準行+1）
COL_QTY = 8            # 数量（基準行+4）
COL_UNIT = 14          # 単位（基準行+4）
COL_NOTE = 17          # 備考（基準行）

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "（1-2）工事内容_雛形.xlsx"


def clear_block(ws, base):
    """1ブロック分の記入セルを空にする（書式・罫線・結合はそのまま）。"""
    for col in LEVEL_COL.values():
        ws.cell(base, col).value = None
    ws.cell(base + 1, COL_SPEC).value = None
    ws.cell(base + 4, COL_QTY).value = None
    ws.cell(base + 4, COL_UNIT).value = None
    ws.cell(base, COL_NOTE).value = None


def write_block(ws, base, item):
    level = item["level"].upper()
    if level not in LEVEL_COL:
        raise ValueError(f"level は B/C/D/E/F/G のいずれか: {item}")
    clear_block(ws, base)
    ws.cell(base, LEVEL_COL[level]).value = item["name"]
    if item.get("spec"):
        ws.cell(base + 1, COL_SPEC).value = item["spec"]
    if item.get("qty") is not None:
        ws.cell(base + 4, COL_QTY).value = item["qty"]
    if item.get("unit"):
        ws.cell(base + 4, COL_UNIT).value = item["unit"]
    if item.get("note"):
        ws.cell(base, COL_NOTE).value = item["note"]


def trim(ws, keep):
    """先頭 keep ブロックだけ残し、以降の行を削除する。"""
    last_row = FIRST_ROW + BLOCK * keep - 1
    for rng in [r for r in ws.merged_cells.ranges if r.min_row > last_row]:
        ws.merged_cells.remove(rng)
    for r in [r for r in ws.row_dimensions if r > last_row]:
        del ws.row_dimensions[r]
    if ws.max_row > last_row:
        ws.delete_rows(last_row + 1, ws.max_row - last_row)
    ws.print_area = f"'{SHEET}'!$B$1:$AB${last_row}"
    return last_row


def extend(ws, need):
    """雛形のブロック数が足りないとき、最後のブロックを複製して増やす。
    書式・罫線・結合・行高をそのまま写すので、様式は変わらない。"""
    from copy import copy as _copy
    cap = (ws.max_row - FIRST_ROW + 1) // BLOCK
    if need <= cap:
        return cap
    src0 = FIRST_ROW + BLOCK * (cap - 1)
    src_merges = [r for r in list(ws.merged_cells.ranges)
                  if src0 <= r.min_row <= src0 + BLOCK - 1]
    ncol = ws.max_column
    for n in range(cap, need):
        dst0 = FIRST_ROW + BLOCK * n
        off = dst0 - src0
        for r in range(BLOCK):
            s, d = src0 + r, dst0 + r
            if ws.row_dimensions[s].height is not None:
                ws.row_dimensions[d].height = ws.row_dimensions[s].height
            for c in range(1, ncol + 1):
                sc = ws.cell(s, c)
                dc = ws.cell(d, c)
                dc.value = None
                dc._style = _copy(sc._style)
        for rng in src_merges:
            ws.merge_cells(start_row=rng.min_row + off, start_column=rng.min_col,
                           end_row=rng.max_row + off, end_column=rng.max_col)
    return need


def build(cfg):
    template = Path(cfg.get("template") or TEMPLATE)
    out = Path(cfg["output"])
    items = cfg["items"]

    wb = openpyxl.load_workbook(template)
    ws = wb[SHEET]

    capacity = (ws.max_row - FIRST_ROW + 1) // BLOCK
    if len(items) > capacity:
        # 雛形より明細が多いときは、最後のブロックを複製して足りない分を作る
        capacity = extend(ws, len(items))

    if cfg.get("title"):
        ws["B1"] = cfg["title"]

    for i, item in enumerate(items):
        write_block(ws, FIRST_ROW + BLOCK * i, item)

    last_row = trim(ws, len(items))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out, len(items), last_row


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    cfg = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out, n, last = build(cfg)
    print(f"出力: {out}")
    print(f"明細 {n} ブロック（B4:AB{last}）")


if __name__ == "__main__":
    main()
