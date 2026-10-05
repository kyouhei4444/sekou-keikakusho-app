# -*- coding: utf-8 -*-
"""「（6）主要資材」様式の Excel を生成する。

usage:  python build_shuyou_shizai.py <入力JSON>

入力JSONの書式は同ディレクトリの README.md を参照。
必要パッケージ: openpyxl
"""
import json
import sys
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "游ゴシック"
NUMFMT = '#,##0.00;"△ "#,##0.00'
WIDTHS = [19.89, 19.89, 9.44, 18.44, 18.44, 9.0, 9.0, 8.89, 9.44]
SHEET_NAME = "(6)　主要資材"

thin = Side(style="thin")

F_TITLE = Font(name=FONT, size=24, bold=True)
F_NOTE = Font(name=FONT, size=9)
F_HEAD = Font(name=FONT, size=12)
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)

AL_LEFT = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
AL_CENTER = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
AL_RIGHT = Alignment(horizontal="right", vertical="bottom", wrap_text=True)
AL_HEAD_D = Alignment(horizontal="distributed", vertical="center")
AL_HEAD_C = Alignment(horizontal="center", vertical="center")
AL_NOTE = Alignment(horizontal="right", vertical="bottom")
AL_TITLE = Alignment(horizontal="centerContinuous", vertical="center")

# 明細列の配置：1=品名, 2=規格, 3=単位, 4=数量, 5=摘要
ALIGN_BY_COL = {1: AL_LEFT, 2: AL_LEFT, 3: AL_CENTER, 4: AL_RIGHT, 5: AL_CENTER}


def _border(top=False, bottom=False, left=True, right=True):
    return Border(left=thin if left else None,
                  right=thin if right else None,
                  top=thin if top else None,
                  bottom=thin if bottom else None)


def build_main_sheet(ws, items, note="「上段赤は実施」"):
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 1行目：標題（選択範囲内で中央 A1:E1）
    ws["A1"] = SHEET_NAME
    for c in range(1, 6):
        ws.cell(row=1, column=c).font = F_TITLE
        ws.cell(row=1, column=c).alignment = AL_TITLE
    ws.row_dimensions[1].height = 39

    # 2行目：注記
    ws["E2"] = note
    ws["E2"].font = F_NOTE
    ws["E2"].alignment = AL_NOTE
    ws["E2"].border = Border(bottom=thin, right=thin)
    ws.row_dimensions[2].height = 16.8

    # 3行目：見出し
    for c, h in enumerate(["品名", "規格", "単位", "数量", "摘要"], start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = F_HEAD
        cell.alignment = AL_HEAD_C if c == 3 else AL_HEAD_D
        cell.border = _border(top=True, bottom=True)
    ws.cell(row=3, column=4).number_format = NUMFMT
    ws.row_dimensions[3].height = 16.8

    # 明細：1明細＝2行。数量は下段（設計）、上段は実施記入用に空欄。
    r = 4
    for it in items:
        top, bot = r, r + 1
        ws.row_dimensions[top].height = 16.8
        ws.row_dimensions[bot].height = 16.8

        for col, val in ((1, it.get("name", "")), (2, it.get("spec", "")),
                         (3, it.get("unit", "")), (5, it.get("remark", ""))):
            ws.merge_cells(start_row=top, start_column=col,
                           end_row=bot, end_column=col)
            ws.cell(row=top, column=col, value=val)
            for rr in (top, bot):
                cell = ws.cell(row=rr, column=col)
                cell.font = F_BODY
                cell.alignment = ALIGN_BY_COL[col]
                cell.border = _border(top=(rr == top), bottom=(rr == bot))

        qty = it.get("qty")
        for rr, val in ((top, None), (bot, None if qty is None else round(float(qty), 5))):
            cell = ws.cell(row=rr, column=4, value=val)
            cell.font = F_BODY
            cell.alignment = AL_RIGHT
            cell.number_format = NUMFMT
            cell.border = _border(top=(rr == top), bottom=(rr == bot))
        r += 2

    ws.freeze_panes = "A4"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if items:
        ws.print_area = "A1:E%d" % (r - 1)
    return r


def build_basis_sheet(ws, header, basis, notes):
    B_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
    AL_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
    AL_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    AL_R = Alignment(horizontal="right", vertical="center")

    for i, w in enumerate([26, 40, 16, 12, 46], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws["A1"] = "主要資材　算出根拠"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)

    r = 2
    for k, v in header.items():
        ws.cell(row=r, column=1, value=k).font = F_BODY
        ws.cell(row=r, column=2, value=v).font = F_BODY
        r += 1
    r += 1

    if basis:
        for i, h in enumerate(["品名", "規格", "算出式", "数量", "根拠資料"], start=1):
            c = ws.cell(row=r, column=i, value=h)
            c.font = F_BOLD
            c.alignment = AL_C
            c.border = B_ALL
        r += 1

        aligns = [AL_L, AL_L, AL_L, AL_R, AL_L]
        for b in basis:
            vals = [b.get("name", ""), b.get("spec", ""), b.get("formula", ""),
                    None if b.get("qty") is None else float(b["qty"]), b.get("source", "")]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=i, value=v)
                c.font = F_BODY
                c.alignment = aligns[i - 1]
                c.border = B_ALL
                if i == 4:
                    c.number_format = "#,##0.00"
            r += 1

    if notes:
        r += 1
        ws.cell(row=r, column=1, value="■ 留意事項").font = F_BOLD
        r += 1
        for i, t in enumerate(notes, start=1):
            c = ws.cell(row=r, column=1, value=str(i))
            c.font = F_BODY
            c.alignment = AL_C
            c2 = ws.cell(row=r, column=2, value=t)
            c2.font = F_BODY
            c2.alignment = AL_L
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            ws.row_dimensions[r].height = 30
            r += 1
    return r


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    with io.open(sys.argv[1], encoding="utf-8") as f:
        cfg = json.load(f)

    out = cfg["output"]
    items = cfg.get("items", [])
    if not items:
        print("ERROR: items が空です")
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    build_main_sheet(ws, items, cfg.get("note", "「上段赤は実施」"))

    if cfg.get("basis") or cfg.get("notes"):
        build_basis_sheet(wb.create_sheet("算出根拠"),
                          cfg.get("header", {}),
                          cfg.get("basis", []),
                          cfg.get("notes", []))

    d = os.path.dirname(os.path.abspath(out))
    if d and not os.path.isdir(d):
        os.makedirs(d)
    wb.save(out)
    print("saved: %s  (%d items)" % (out, len(items)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
