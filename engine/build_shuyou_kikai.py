# -*- coding: utf-8 -*-
"""「（5）主要機械」様式の Excel を生成する。

usage:  python build_shuyou_kikai.py <入力JSON>

入力JSONの書式は同ディレクトリの README.md を参照。
必要パッケージ: openpyxl
"""
import json
import sys
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

FONT = "Yu Gothic UI"
SHEET_NAME = "主要機械"
BASIS_SHEET = "算出根拠・留意事項"

# 列: A=左余白, B=機械名, C=規格, D=単位, E=数量, F=工種, G=低騒音, H=排ガス
# A4縦（印刷可能幅 約180mm）に収まる幅。合計は概ね 90 前後に抑える。
WIDTHS = {"A": 2, "B": 22, "C": 22, "D": 5, "E": 5, "F": 22, "G": 6.5, "H": 6.5}
HEADERS = ["", "機　械　名", "　規　格", "単 位", "数 量", "工　種", "低騒音", "排ガス"]
# 明細フィールドの並び（B列以降）
FIELDS = ["name", "spec", "unit", "qty", "work", "low_noise", "emission"]
# 左寄せにする列インデックス（0=name, 1=spec, 4=work）。他は中央。
LEFT_COLS = {0, 1, 4}

thin = Side(style="thin", color="000000")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
AL_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")

F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True)
F_BODY = Font(name=FONT, size=9)


def _set_a4_portrait(ws):
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.6
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3


def build_main_sheet(ws, machines, title="（5）主要機械",
                     note="低騒音、排出ガス対策型の建設機械を使用する。"):
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    # 1行目：標題
    ws["B1"] = title
    ws["B1"].font = F_TITLE
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")

    # 3行目：注記
    if note:
        ws["C3"] = note
        ws["C3"].font = F_BODY
        ws["C3"].alignment = Alignment(horizontal="left", vertical="center")

    # 5行目：見出し
    hrow = 5
    for i, h in enumerate(HEADERS):
        c = ws.cell(row=hrow, column=i + 1, value=h)
        if i >= 1:
            c.font = F_HEAD
            c.alignment = AL_CENTER
            c.border = BOX
            c.fill = HEAD_FILL

    # 明細（B列以降）
    r = hrow + 1
    for m in machines:
        for i, key in enumerate(FIELDS):
            val = m.get(key, "")
            c = ws.cell(row=r, column=i + 2, value=val)
            c.font = F_BODY
            c.border = BOX
            c.alignment = AL_LEFT if i in LEFT_COLS else AL_CENTER
        ws.row_dimensions[r].height = 28
        r += 1
    last = r - 1

    ws.sheet_view.showGridLines = False
    _set_a4_portrait(ws)
    if machines:
        ws.print_area = "A1:H%d" % last
    return last


def build_basis_sheet(ws, notes):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 92
    for i, line in enumerate(notes, start=1):
        bold = (i == 1) or str(line).startswith("■")
        c = ws.cell(row=i, column=2, value=line)
        c.font = Font(name=FONT, size=10, bold=bold)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    _set_a4_portrait(ws)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    with io.open(sys.argv[1], encoding="utf-8") as f:
        cfg = json.load(f)

    out = cfg["output"]
    machines = cfg.get("machines", [])
    if not machines:
        print("ERROR: machines が空です")
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    build_main_sheet(ws, machines,
                     cfg.get("title", "（5）主要機械"),
                     cfg.get("note", "低騒音、排出ガス対策型の建設機械を使用する。"))

    if cfg.get("notes"):
        build_basis_sheet(wb.create_sheet(BASIS_SHEET), cfg["notes"])

    d = os.path.dirname(os.path.abspath(out))
    if d and not os.path.isdir(d):
        os.makedirs(d)
    wb.save(out)
    print("saved: %s  (%d machines)" % (out, len(machines)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
