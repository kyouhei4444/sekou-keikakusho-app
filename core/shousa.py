# -*- coding: utf-8 -*-
"""設計照査：取り込んだ設計図書どうしを突き合わせて、食い違い・記載漏れを拾う。

根拠
  土木工事共通仕様書［令和7年10月］1-1-1-3「設計図書の照査」2.（p.1-5）
    受注者は施工前及び施工途中に、契約書第18条第1項第1号から第5号に係る
    設計図書の照査を行い、該当する事実があれば監督員に資料を提出して確認を求める。
  建設工事請負契約約款 第18条第1項（条件変更等）
    一 図面、仕様書、現場説明書及び質問回答書が一致しないこと
    二 設計図書に誤謬又は脱漏があること
    三 設計図書の表示が明確でないこと
    四 施工条件と実際の工事現場が一致しないこと
    五 設計図書で明示されていない施工条件について予期できない特別な状態が生じたこと

アプリが**機械的に確かめられるのは第1号〜第3号の一部**（資料の文字どうしの照合）だけ。
第4号・第5号（現地と図面の照合）と図面の中身は、照査チェックリストとして人が確かめる。
ここで出す指摘は「要確認」の候補であり、誤りと決めつけない（読み取りの誤りも混ざる）。
"""
import datetime
import difflib
import os
import re
import unicodedata

from . import meisai, soukatsu, util

try:
    import pymupdf
except Exception:                       # pragma: no cover
    pymupdf = None

GOU = {
    0: '（照合できなかった資料）',
    1: '第1号 図書の不一致',
    2: '第2号 誤謬・脱漏',
    3: '第3号 表示が不明確',
    4: '第4号 現場条件の相違',
    5: '第5号 予期できない特別な状態',
}
KONKYO = ('土木工事共通仕様書［令和7年10月］1-1-1-3 設計図書の照査（p.1-5）／'
          '建設工事請負契約約款 第18条第1項')

# 図書間で比べるときの基準にする順
BASE_ORDER = ['工事数量総括表', '見積参考資料', '公表用設計書', '設計書']
ROUND_BASIS = (1, 10, 100, 1000)        # 明細書の「○○当り」によく使われる数
SMALL_QTY = 0.05                        # これ以下の数量は表示を確かめる


def norm(s):
    return unicodedata.normalize('NFKC', str(s or '')).replace(' ', '').replace('　', '')


def _unit(u):
    return norm(u).replace('㎡', 'm2').replace('㎥', 'm3').lower()


def _num(v):
    if v in (None, ''):
        return None
    try:
        f = float(str(v).replace(',', ''))
    except ValueError:
        return None
    return int(f) if f == int(f) else f


def _same(a, b):
    return abs(a - b) <= max(1e-6, 1e-4 * max(abs(a), abs(b)))


def _fmt(v):
    if v is None:
        return '（なし）'
    if isinstance(v, float):
        v = round(v, 4)
        if v == int(v):
            v = int(v)
    return '{:,}'.format(v) if isinstance(v, int) else str(v)


def _q(v, u):
    return '%s%s' % (_fmt(v), u or '')


class Findings(object):
    def __init__(self):
        self.items = []

    def add(self, gou, level, target, text, source='', page=None, kind=''):
        src = source
        if page:
            src = '%s p.%s' % (source, page) if source else 'p.%s' % page
        self.items.append({'gou': gou, '区分': GOU[gou], '重要度': level, '対象': target,
                           '内容': text, '資料': src, '種類': kind})


# ---------------------------------------------------------------- 本体
def run(files, project=None, log=None):
    """files … 判別結果（path／kind／内訳あり／単価表あり）。project … 工事情報。"""
    log = log or (lambda s: None)
    project = project or {}
    if pymupdf is None:
        return {'error': 'pymupdf が入っていません'}
    files = [dict(f) for f in (files or []) if f.get('path') and os.path.isfile(f['path'])]
    if not files:
        return {'error': '照査する資料がありません。先に「判別する」を押してください'}
    _names(files)

    fd = Findings()
    checked = []                # 実施した照合の一覧（何を見て、何を見ていないか）

    # 1) 内訳書の読み取り（工事数量総括表・公表用設計書・見積参考資料）
    docs = []
    for f in files:
        if not f.get('内訳あり'):
            continue
        log('内訳書を読んでいます：%s' % f['name'])
        r = soukatsu.parse_naiyou(f['path'])
        if r.get('error') or not r.get('rows'):
            fd.add(0, '参考', f['name'], '内訳書の表を読み取れませんでした（画像PDF、または内訳書'
                   'ではない資料）。必要なら目視で照合してください', f['name'], kind='読み取り')
            continue
        docs.append({'file': f, 'rows': r['rows'], 'meta': r.get('meta') or {}})
    # 基準は直下のフォルダーの資料を優先（サブフォルダーには旧版・変更版が入りがち）
    docs.sort(key=lambda d: (d['file']['name'].count('/'),
                             BASE_ORDER.index(d['file']['kind'])
                             if d['file']['kind'] in BASE_ORDER else 9, d['file']['name']))

    compare = []
    if len(docs) >= 2:
        base, diffs = docs[0], []
        for other in docs[1:]:
            compare += _compare_docs(base, other, fd, diffs)
        _report_diffs(diffs, base['file']['name'], fd)
        checked.append('図書間の内訳書（工種・数量・単位）の照合：%s' %
                       '／'.join(d['file']['name'] for d in docs))
    elif docs:
        fd.add(1, '参考', '図書間の照合', '内訳書のある資料が1つ（%s）だけなので、図書間の数量の照合は'
               'していません' % docs[0]['file']['name'], kind='図書間')

    if docs:
        _check_rows(docs[0], fd)
        checked.append('内訳書の数量・単位の記載漏れ、極端に小さい数量：%s' % docs[0]['file']['name'])

    # 2) 工事名・工事場所の照合
    metas = [(d['file']['name'], d['meta']) for d in docs]
    for f in files:
        if f.get('内訳あり') or f.get('kind') not in ('特記仕様書', '設計書', '数量計算書'):
            continue
        try:
            doc = pymupdf.open(f['path'])
            metas.append((f['name'], soukatsu.read_meta(doc)))
            doc.close()
        except Exception:
            pass
    _check_meta(metas, project, fd)
    checked.append('工事名・工事場所の照合（各資料の表紙と、工事情報タブの入力）')

    # 3) 見積参考資料の明細書（内訳書 → 工種明細書 → 科目明細書）
    for f in files:
        if not (f.get('内訳あり') and f.get('単価表あり')):
            continue
        log('明細書をたどっています：%s' % f['name'])
        try:
            book = meisai.read_book(f['path'])
        except Exception as e:
            fd.add(0, '参考', f['name'], '明細書を読み取れませんでした（%s）' % e, f['name'],
                   kind='読み取り')
            continue
        if not book['blocks']:
            continue
        _check_book(book, f['name'], fd, _ref_counter(f['path']))
        checked.append('内訳書と明細書の照合（参照先の有無・単位・数量）：%s' % f['name'])
        break

    # 4) 特記仕様書の工期
    for f in files:
        if f.get('kind') != '特記仕様書':
            continue
        log('特記仕様書を読んでいます：%s' % f['name'])
        _check_kouki(f, project, fd)
        checked.append('特記仕様書の工期と、工事情報タブの工期の照合：%s' % f['name'])
        break

    order = {'要確認': 0, '参考': 1}
    fd.items.sort(key=lambda x: (order.get(x['重要度'], 2), x['gou']))
    for i, x in enumerate(fd.items):
        x['no'] = i + 1
    n_ng = sum(1 for x in fd.items if x['重要度'] == '要確認')
    return {
        'findings': fd.items, 'compare': compare, 'checked': checked,
        'docs': [{'name': f['name'], 'kind': f.get('kind', ''), 'pages': f.get('pages', '')}
                 for f in files],
        'checklist': CHECKLIST, 'date': datetime.date.today().isoformat(),
        'summary': '要確認 %d 件／参考 %d 件' % (n_ng, len(fd.items) - n_ng),
        'konkyo': KONKYO,
    }


def _names(files):
    """表示名。同じファイル名が別のフォルダーにあるときは、フォルダー名を付けて区別する。"""
    try:
        root = os.path.commonpath([os.path.dirname(f['path']) for f in files])
    except ValueError:
        root = ''
    for f in files:
        rel = os.path.relpath(f['path'], root) if root else os.path.basename(f['path'])
        f['name'] = rel.replace(os.sep, '/')


# ---------------------------------------------------------------- 図書間の内訳書
def _key(r):
    return norm(r.get('name')) + '|' + norm(r.get('spec'))


def _label(r):
    return (r.get('name') or '') + ((' ' + r['spec']) if r.get('spec') else '')


def _countable(r):
    """数量の照合に使える行（集計行・一式は読み取りが崩れやすいので外す）。"""
    return r.get('qty') is not None and _unit(r.get('unit')) not in ('', '式')


def _compare_docs(base, other, fd, diffs):
    a, b = base['rows'], other['rows']
    an, bn = base['file']['name'], other['file']['name']
    sm = difflib.SequenceMatcher(None, [_key(r) for r in a], [_key(r) for r in b], autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == 'equal':
            for ra, rb in zip(a[i1:i2], b[j1:j2]):
                qa, qb = _num(ra.get('qty')), _num(rb.get('qty'))
                ua, ub = _unit(ra.get('unit')), _unit(rb.get('unit'))
                judge = ''
                if qa is not None and qb is not None and not _same(qa, qb):
                    judge = '数量が違う'
                elif ua and ub and ua != ub:
                    judge = '単位が違う'
                if _countable(ra) or _countable(rb) or judge:
                    out.append({'name': _label(ra), 'base': an, 'other': bn,
                                'a': _q(ra.get('qty'), ra.get('unit')),
                                'b': _q(rb.get('qty'), rb.get('unit')),
                                'pa': ra.get('page'), 'pb': rb.get('page'), 'judge': judge or '一致'})
                if judge:
                    diffs.append({'row': ra, 'judge': judge, 'doc': bn, 'other': rb})
            continue
        # 片方にしか無い行。一式・集計行は読み取りの崩れが多いので、数量のある行だけ挙げる
        for side, rows, mine, theirs in (('a', a[i1:i2], an, bn), ('b', b[j1:j2], bn, an)):
            for r in rows:
                if not _countable(r):
                    continue
                qtxt = _q(r.get('qty'), r.get('unit'))
                out.append({'name': _label(r), 'base': an, 'other': bn,
                            'a': qtxt if side == 'a' else '（なし）',
                            'b': qtxt if side == 'b' else '（なし）',
                            'pa': r.get('page') if side == 'a' else None,
                            'pb': r.get('page') if side == 'b' else None, 'judge': '片方にない'})
                fd.add(1, '要確認', _label(r), '%s にある明細（%s）が %s に見当たりません。'
                       '名称の表記違い・読み取りの誤りの可能性もあります' % (mine, qtxt, theirs),
                       mine, r.get('page'), kind='図書間')
    return out


def _report_diffs(diffs, an, fd):
    """同じ明細の食い違いは1件にまとめる（基準の資料と、違っていた資料を並べる）。"""
    groups = {}
    for d in diffs:
        groups.setdefault((id(d['row']), d['judge']), []).append(d)
    for (_, judge), ds in groups.items():
        ra = ds[0]['row']
        vals = ['%s %s' % (an, _q(ra.get('qty'), ra.get('unit')))]
        pages = ['%s p.%s' % (an, ra.get('page') or '?')]
        for d in ds:
            rb = d['other']
            vals.append('%s %s' % (d['doc'], _q(rb.get('qty'), rb.get('unit'))))
            pages.append('%s p.%s' % (d['doc'], rb.get('page') or '?'))
        fd.add(1, '要確認', _label(ra), '%s：%s' % (judge, '、'.join(vals)), '／'.join(pages),
               kind='図書間')


def _check_rows(doc, fd):
    """基準にした内訳書の、細別の数量・単位の漏れと、極端に小さい数量。"""
    name = doc['file']['name']
    rows = doc['rows']
    for i, r in enumerate(rows):
        if r.get('level') != 'G':
            continue
        q, u = _num(r.get('qty')), (r.get('unit') or '').strip()
        nxt = rows[i + 1] if i + 1 < len(rows) else None
        if (q is None and nxt and nxt.get('level') == 'G' and nxt.get('qty') is not None
                and nxt.get('page') == r.get('page')):
            continue      # 名称と規格が2行に分かれ、数量が下の行に印字される様式（記載漏れではない）
        if q is None or not u:
            what = '数量・単位' if (q is None and not u) else ('数量' if q is None else '単位')
            fd.add(2, '要確認', _label(r), '細別に%sの記載がありません' % what,
                   name, r.get('page'), kind='内訳書')
        elif q == 0:
            fd.add(2, '要確認', _label(r), '数量が 0 になっています', name, r.get('page'),
                   kind='内訳書')
        elif q <= SMALL_QTY and _unit(u) != '式':
            fd.add(3, '参考', _label(r), '数量が極端に小さい（%s）。計上の趣旨（仮計上・概算など）を'
                   '監督員に確かめてください' % _q(q, u), name, r.get('page'), kind='内訳書')


# ---------------------------------------------------------------- 工事名・工事場所
def _clean_meta(v, labels):
    t = unicodedata.normalize('NFKC', str(v or ''))
    t = re.sub(r'^[（(]当初[）)]', '', t).strip()
    for lab in labels:                    # 「工事場所名 …」のようにラベルが混ざる様式
        t = re.sub('^' + re.escape(lab) + r'名?\s*', '', t)
    toks = []
    for x in t.split():
        if x not in toks:
            toks.append(x)                # 同じ値が2回印字される様式
    return ''.join(toks)


def _check_meta(metas, project, fd):
    for key, labels in (('工事名', ['工事名', '工事件名']),
                        ('工事場所', ['工事場所', '施工場所', '工事箇所'])):
        vals = []
        for src, m in metas:
            v = _clean_meta(m.get(key), labels)
            if len(v) >= 3:
                vals.append((src, v))
        inp = _clean_meta(project.get(key), labels)
        if inp:
            vals.append(('工事情報タブの入力', inp))
        if len(vals) < 2:
            continue
        ref_src, ref = vals[0]
        for src, v in vals[1:]:
            a, b = norm(ref), norm(v)
            if a in b or b in a:
                continue
            fd.add(1, '要確認', key, '%s が一致しません：%s「%s」／%s「%s」' % (key, ref_src, ref, src, v),
                   '%s／%s' % (ref_src, src), kind='表紙')


# ---------------------------------------------------------------- 明細書
def _kei(block):
    """工種明細書の「計」行（何単位あたりか）。"""
    for r in block['items']:
        if norm(meisai.rec_name(r)).startswith('計'):
            return r
    return None


def _label_block(bid):
    kind, no = bid[:2], bid[2:]
    return {'工種': '第%s号工種明細書', '科目': '第%s号科目明細書',
            '単価': '第0-%s号施工単価表'}.get(kind, '%s') % no


def _ref_counter(pdf):
    """参照先の見出しが本文に出てくる回数を数える関数を返す。
    表の読み取りで取りこぼした明細書を「参照切れ」と誤って挙げないための確認。"""
    try:
        doc = pymupdf.open(pdf)
        text = ''.join(norm(doc[i].get_text('text')) for i in range(doc.page_count))
        doc.close()
    except Exception:
        return lambda bid: 99

    def count(bid):
        kind, no = bid[:2], bid[2:]
        if kind == '単価':
            pat = r'第0-0*%d号施工単価' % int(no)
        else:
            pat = r'第0*%d号%s明細書' % (int(no), kind)
        return len(re.findall(pat, text))
    return count


def _check_book(book, src, fd, count=None):
    blocks = book['blocks']
    tanka = book.get('tanka') or {}
    count = count or (lambda bid: 0)
    missing = set()
    for r in book['naiyou']:
        ref = r.get('ref')
        if not ref or not ref.startswith(('工種', '科目')):
            continue
        name = meisai.rec_name(r)
        blk = blocks.get(ref)
        if blk is None:
            if ref not in missing and count(ref) < 2:      # 備考の参照と見出しで2回は出る
                missing.add(ref)
                fd.add(2, '要確認', name, '内訳書の備考にある %s が資料の中に見当たりません'
                       % _label_block(ref), src, r.get('page'), kind='明細書')
            continue
        q, u = _num(r.get('qty')), r.get('unit') or ''
        if blk['kind'] == '工種':
            kei = _kei(blk)
            if kei and kei.get('unit') and u and _unit(kei['unit']) != _unit(u):
                fd.add(1, '要確認', name, '単位が違う：内訳書は %s、%s の計は %s'
                       % (u, _label_block(ref), _q(kei.get('qty'), kei['unit'])),
                       '%s p.%s／p.%s' % (src, r.get('page'), blk.get('page')), kind='明細書')
            basis = blk.get('basis') or 1
            for it in blk['items']:
                sub = it.get('ref')
                if not sub or not sub.startswith(('工種', '科目')):
                    continue
                child = blocks.get(sub)
                if child is None:
                    if sub not in missing and count(sub) < 2:
                        missing.add(sub)
                        fd.add(2, '要確認', name, '%s の備考にある %s が資料の中に見当たりません'
                               % (_label_block(ref), _label_block(sub)), src, blk.get('page'),
                               kind='明細書')
                    continue
                if q is not None and it.get('qty') is not None:
                    _check_heading(name, child, sub, q * it['qty'] / basis, it.get('unit'),
                                   src, blk.get('page'), fd)
        elif blk['kind'] == '科目' and q is not None:
            _check_heading(name, blk, ref, q, u, src, r.get('page'), fd)

    # 明細書から施工単価表への参照（単価表の読み取りが崩れることがあるので「参考」）
    for bid, blk in blocks.items():
        if blk['kind'] == '単価':
            continue
        for it in blk['items']:
            ref = it.get('ref')
            if (not ref or not ref.startswith('単価') or ref in blocks
                    or ref[2:].zfill(4) in tanka or ref in missing or count(ref) >= 2):
                continue
            missing.add(ref)
            fd.add(2, '参考', meisai.rec_name(it) or blk['name'],
                   '%s の備考にある %s が見つかりませんでした（資料に無いか、読み取れていない）'
                   % (_label_block(bid), _label_block(ref)), src, it.get('page'), kind='明細書')


def _check_heading(name, child, cid, need, unit, src, page, fd):
    """科目明細書の見出し数量が、内訳書から計算した数量と合っているか。
    見出しが 1・10・100・1000 のときは「○○当り」の単価なので比べない。"""
    hq = _num(child.get('qty'))
    if hq is None:
        return
    if unit and child.get('unit') and _unit(unit) != _unit(child['unit']):
        fd.add(1, '要確認', name, '単位が違う：上位の明細は %s、%s の見出しは %s'
               % (unit, _label_block(cid), _q(hq, child['unit'])),
               '%s p.%s／p.%s' % (src, page, child.get('page')), kind='明細書')
        return
    if hq in ROUND_BASIS or _same(hq, need):
        return
    fd.add(1, '要確認', name, '数量が違う：内訳書から計算すると %s、%s の見出しは %s'
           % (_q(round(need, 3), unit), _label_block(cid), _q(hq, child.get('unit'))),
           '%s p.%s／p.%s' % (src, page, child.get('page')), kind='明細書')


# ---------------------------------------------------------------- 特記仕様書の工期
DATE = r'令和\s*(\d{1,2}|元)\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日'


def _to_date(m):
    y = 1 if m.group(1) == '元' else int(m.group(1))
    try:
        return datetime.date(2018 + y, int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def read_kouki(pdf):
    """特記仕様書から工期を読む。
    「工期：令和8年10月8日から令和9年6月30日まで」（県）と
    「工期は、契約日の翌日から令和9年2月26日とする」（市）の2通り。"""
    doc = pymupdf.open(pdf)
    try:
        for pi in range(min(doc.page_count, 40)):
            t = unicodedata.normalize('NFKC', doc[pi].get_text('text'))
            t = re.sub(r'\s+', ' ', t)
            for m in re.finditer(r'工 ?期 ?(?::|は)', t):
                win = t[m.end():m.end() + 90]
                cut = re.search(r'(まで|とする)', win)
                if cut:
                    win = win[:cut.end()]
                ds = list(re.finditer(DATE, win))
                if not ds:
                    continue
                end = _to_date(ds[-1])
                start = _to_date(ds[0]) if len(ds) >= 2 else None
                how = '' if start else win[:ds[-1].start()].strip(' 、,から')
                return {'自': start, '至': end, '自の書き方': how, 'page': pi + 1,
                        'text': ('工期：' + win.strip(' 、,')).strip()}
    finally:
        doc.close()
    return None


def _check_kouki(f, project, fd):
    try:
        k = read_kouki(f['path'])
    except Exception:
        k = None
    if not k or not k.get('至'):
        fd.add(3, '参考', '工期', '特記仕様書から工期を読み取れませんでした'
               '（「契約書に記載の期間」とする書き方の場合もあります）。契約書で確かめてください',
               f['name'], kind='工期')
        return
    jsi, jitaru = util.parse_date(project.get('工期自')), util.parse_date(project.get('工期至'))
    shown = '%s～%s' % (util.wareki(k['自']) if k['自'] else (k['自の書き方'] or '（始期の記載なし）'),
                       util.wareki(k['至']))
    if not jitaru:
        fd.add(1, '参考', '工期', '特記仕様書の工期は %s。工事情報タブに工期が入っていません' % shown,
               f['name'], k['page'], kind='工期')
        return
    ng = []
    if k['至'] != jitaru:
        ng.append('終期（特記仕様書 %s／入力 %s）' % (util.wareki(k['至']), util.wareki(jitaru)))
    if k['自'] and jsi and k['自'] != jsi:
        ng.append('始期（特記仕様書 %s／入力 %s）' % (util.wareki(k['自']), util.wareki(jsi)))
    if ng:
        fd.add(1, '要確認', '工期', '工期が一致しません：%s。余裕期間・契約書の工期も確かめてください'
               % '、'.join(ng), f['name'], k['page'], kind='工期')


# ---------------------------------------------------------------- 照査チェックリスト（人が確かめる）
# 区分は契約約款第18条第1項の各号。項目はアプリが用意した確認の目安で、基準の条文そのものではない。
CHECKLIST = [
    (1, '図面と数量総括表', '図面（平面・縦断・横断・構造図）の数量・寸法と、工事数量総括表・数量計算書の数量が一致しているか'),
    (1, '図面相互', '平面図・縦断図・横断図・構造図・展開図の間で、延長・高さ・幅・測点が一致しているか'),
    (1, '特記仕様書と図面', '特記仕様書の材料規格（コンクリート強度・鉄筋・製品規格）と、図面・数量総括表の規格が一致しているか'),
    (1, '現場説明書・質問回答書', '入札時の質問回答・現場説明の内容が、図面・特記仕様書に反映されているか'),
    (1, '数量計算書', '数量計算書の集計値と、工事数量総括表の数量が一致しているか'),
    (2, '必要な図面の有無', '施工に必要な図面（構造図・配筋図・標準図・詳細図・仮設図）が揃っているか'),
    (2, '数量の計上漏れ', '図面にある構造物・付帯工（目地・水抜き・基礎・小口止め等）が数量総括表に計上されているか'),
    (2, '仮設・準備工', '工事用道路・仮締切・排水・交通規制など、現場条件から必要な仮設が計上されているか'),
    (2, '計算の誤り', '数量計算書の寸法の拾い・計算式・単位換算に誤りがないか'),
    (3, '寸法・位置の表示', '寸法・標高・測点・施工範囲（起終点）が図面で読み取れるか'),
    (3, '材料・規格の表示', '使用材料の規格・品質・製品の型式が特定できるか'),
    (3, '施工方法の指定', '指定仮設・施工方法の指定の有無と、その内容が明確か'),
    (4, '地形', '現地の地形（地盤高・横断形状）が図面と一致しているか（着手前測量で確認）'),
    (4, '地質・湧水', '土質・岩の出現・湧水・地下水位が、設計の前提と一致しているか'),
    (4, '地下埋設物・架空線', '水道・ガス・電力・通信の埋設物、架空線の位置が把握され、図面と一致しているか'),
    (4, '用地・境界', '用地境界・借地範囲・作業ヤードが確保でき、図面と一致しているか'),
    (4, '近接構造物・家屋', '近接する構造物・家屋・井戸への影響と、事前調査の要否'),
    (4, '交通・搬入路', '交通規制の条件、資材搬入路・大型車の進入可否が設計の前提と一致しているか'),
    (4, '河川・水路の条件', '河川内・水路の工事では、水位・流量・渇水期・水利権者との調整条件'),
    (5, '想定外の条件', '設計図書に示されていない支障物（転石・埋設物・旧構造物等）や特殊な状態の有無と、見つけたときの記録・報告方法'),
]


# ---------------------------------------------------------------- Excel 出力
def make_xlsx(prj, res, dst_dir, log=print):
    """設計照査結果.xlsx を作る（照査結果／図書間の数量比較／照査チェックリスト／照査した資料）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    if not dst_dir:
        raise RuntimeError('出力先フォルダーが指定されていません')
    os.makedirs(dst_dir, exist_ok=True)
    out = os.path.join(dst_dir, '設計照査結果.xlsx')

    thin = Side(style='thin', color='808080')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor='DDE4EE')
    ng_fill = PatternFill('solid', fgColor='FCE8E6')
    font = Font(name='ＭＳ ゴシック', size=10)
    bold = Font(name='ＭＳ ゴシック', size=10, bold=True)
    title = Font(name='ＭＳ ゴシック', size=14, bold=True)
    wrap = Alignment(wrap_text=True, vertical='top')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def setup(ws, widths, landscape=True, head_row=None):
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.6
        ws.oddFooter.center.text = '&P / &N'
        if head_row:
            ws.print_title_rows = '%d:%d' % (head_row, head_row)
            ws.freeze_panes = 'A%d' % (head_row + 1)

    def header(ws, r, labels):
        for c, t in enumerate(labels, 1):
            x = ws.cell(r, c, t)
            x.font, x.fill, x.border, x.alignment = bold, head_fill, box, center

    def row(ws, r, vals, fill=None, centers=()):
        lines = 1
        for c, v in enumerate(vals, 1):
            x = ws.cell(r, c, v)
            x.font, x.border = font, box
            x.alignment = center if c in centers else wrap
            if fill:
                x.fill = fill
            # 折り返す行は高さを自動で広げてくれないので、文字幅から行数を見積もる
            w = ws.column_dimensions[chr(64 + c)].width or 8.43
            n = sum(2 if ord(ch) > 0xff else 1 for ch in str(v if v is not None else ''))
            lines = max(lines, -(-n // max(1, int(w * 0.95))))
        ws.row_dimensions[r].height = 13.5 * lines + 3

    wb = Workbook()
    koji = prj.get('工事名') or ''
    today = util.parse_date(res.get('date')) or datetime.date.today()

    # ---- 照査結果
    ws = wb.active
    ws.title = '照査結果'
    setup(ws, [8, 20, 8, 26, 56, 30, 22], head_row=7)
    ws['A1'] = '設計照査結果（資料の照合）'
    ws['A1'].font = title
    items = [x for x in res.get('findings') or [] if not x.get('除')]
    n_ng = sum(1 for x in items if x.get('重要度') == '要確認')
    info = [('工事名', koji), ('照査日', util.wareki(today)),
            ('根拠', res.get('konkyo') or KONKYO),
            ('結果', '要確認 %d 件／参考 %d 件　※アプリによる資料の文字の照合。読み取りの誤りも'
                    '含むので、元の資料で確かめてから監督員に確認を求めること' % (n_ng, len(items) - n_ng))]
    for i, (k, v) in enumerate(info):
        ws.cell(2 + i, 1, k).font = bold
        ws.cell(2 + i, 2, v).font = font
        ws.merge_cells(start_row=2 + i, start_column=2, end_row=2 + i, end_column=7)
        ws.cell(2 + i, 2).alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[5].height = 28
    header(ws, 7, ['No', '照査区分（約款第18条第1項）', '重要度', '対象', '内容', '資料・ページ',
                   '対応（監督員への確認・回答）'])
    r = 8
    for i, x in enumerate(items, 1):
        row(ws, r, [i, x.get('区分'), x.get('重要度'), x.get('対象'), x.get('内容'), x.get('資料'), ''],
            ng_fill if x.get('重要度') == '要確認' else None, centers=(1, 3))
        r += 1
    if not items:
        row(ws, r, ['', '', '', '', '資料の照合では指摘事項はありませんでした', '', ''])
        r += 1
    r += 1
    ws.cell(r, 1, '実施した照合').font = bold
    for t in res.get('checked') or []:
        r += 1
        ws.cell(r, 2, '・' + t).font = font
    r += 1
    ws.cell(r, 2, '・現地条件（第4号・第5号）と図面の中身は「照査チェックリスト」シートで確かめる').font = font

    # ---- 図書間の数量比較
    cmp_ = res.get('compare') or []
    ws = wb.create_sheet('図書間の数量比較')
    setup(ws, [5, 46, 24, 14, 24, 14, 12], head_row=3)
    ws['A1'] = '図書間の内訳書の数量比較'
    ws['A1'].font = title
    header(ws, 3, ['No', '名称・規格', '資料A', '数量A', '資料B', '数量B', '判定'])
    for i, x in enumerate(cmp_, 1):
        row(ws, 3 + i, [i, x['name'], '%s p.%s' % (x['base'], x.get('pa') or '-'), x['a'],
                        '%s p.%s' % (x['other'], x.get('pb') or '-'), x['b'], x['judge']],
            ng_fill if x['judge'] != '一致' else None, centers=(1, 7))
    if not cmp_:
        ws.cell(4, 2, '比べられる資料が1つしかないため、比較していません').font = font

    # ---- 照査チェックリスト
    ws = wb.create_sheet('照査チェックリスト')
    setup(ws, [5, 23, 20, 60, 8, 11, 11, 22], head_row=4)
    ws['A1'] = '設計照査チェックリスト（現地・図面）'
    ws['A1'].font = title
    ws['A2'] = ('区分は建設工事請負契約約款 第18条第1項の各号。確認項目は照査の目安（アプリ作成）であり、'
                '工事の内容に合わせて追加・削除すること。結果欄：○該当なし／×該当あり（監督員に確認）／－対象外')
    ws['A2'].font = font
    ws.merge_cells('A2:H2')
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 30
    header(ws, 4, ['No', '照査区分', '確認項目', '確認の観点', '結果', '確認日', '確認者', '備考'])
    dv = DataValidation(type='list', formula1='"○,×,－"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (g, item, point) in enumerate(res.get('checklist') or CHECKLIST, 1):
        row(ws, 4 + i, [i, GOU[g], item, point, '', '', '', ''], centers=(1, 5))
        dv.add('E%d' % (4 + i))

    # ---- 照査した資料
    ws = wb.create_sheet('照査した資料')
    setup(ws, [5, 60, 18, 10], landscape=False, head_row=3)
    ws['A1'] = '照査した資料'
    ws['A1'].font = title
    header(ws, 3, ['No', 'ファイル', '種類', 'ページ'])
    for i, d in enumerate(res.get('docs') or [], 1):
        row(ws, 3 + i, [i, d.get('name'), d.get('kind'), d.get('pages')], centers=(1, 4))

    wb.save(out)
    log('設計照査結果.xlsx  指摘 %d 件／チェック項目 %d 件'
        % (len(items), len(res.get('checklist') or CHECKLIST)))
    return out
